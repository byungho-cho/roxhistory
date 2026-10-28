"""Build the Word edition from the authoritative Markdown synopsis."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
FONT = 'Noto Sans CJK KR'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
section.header_distance = section.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = FONT
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_after = Pt(7)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing = 1.28
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
header = section.header.paragraphs[0]
header.text = 'RoxHistory   초기 기획 기준 문서'
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('v0.1   2026년 10월 28일     ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)

page_starts = {'세 가지 시작 경로', '장 진행과 시대 전환', '역사와 지리가 남는 방식', '비주얼과 웹 구현 방향', '다음 디자인 작업'}
lines = (ROOT/'docs/synopsis-v0.1.md').read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        heading=line[3:]
        p=doc.add_paragraph(heading,'Heading 1')
        if heading in page_starts:p.paragraph_format.page_break_before=True
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',v) for v in vals):rows.append(vals)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]));table.autofit=False
        widths=([1.02,2.0,1.58,2.4] if len(rows[0])==4 else [4.7,2.3])
        for col,w in zip(table.columns,widths):col.width=Inches(w)
        for ri,row in enumerate(rows):
            cells=table.add_row().cells
            for ci,(cell,value) in enumerate(zip(cells,row)):
                cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=value
                tcPr=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263D4B' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'5');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                tcPr.append(borders)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'95');edge.set(qn('w:type'),'dxa');margins.append(edge)
                tcPr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.15
                    for r in p.runs:r.font.size=Pt(9);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
            if ri==0:
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
        continue
    else:
        p=doc.add_paragraph(line[2:] if line.startswith('- ') else line)
        if line.startswith('- '):
            p.paragraph_format.left_indent=Inches(.12)
        if 'https://' in line:
            p.paragraph_format.space_after=Pt(6)
            for r in p.runs:r.font.size=Pt(8)
    i+=1
doc.core_properties.title='RoxHistory 초기 시놉시스'
doc.core_properties.subject='게임 콘셉트와 합의 범위'
doc.core_properties.author='RoxHistory'
for element in doc.styles.element.iter():
    for child in list(element):
        if child.tag == qn('w:pBdr'):element.remove(child)
for element in doc.element.iter():
    for child in list(element):
        if child.tag == qn('w:pBdr'):element.remove(child)
doc.save(ROOT/'docs/RoxHistory_Synopsis_v0.1.docx')
print(ROOT/'docs/RoxHistory_Synopsis_v0.1.docx')
